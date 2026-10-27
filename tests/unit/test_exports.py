"""Unit tests for the report exports (T9).

The fixtures are deliberately built from the real T2/T4 artifacts where a value
matters (125 municipalities, null insufficient rows, model/version metadata) and
from synthetic rows only for the null/formula edge cases.
"""

from __future__ import annotations

import csv
import io
import json
import re
from dataclasses import replace
from datetime import UTC, datetime
from pathlib import Path
from uuid import UUID

import pytest
from openpyxl import load_workbook
from PIL import Image
from pypdf import PdfReader

from participa_geo.api.schemas.common import DatasetMeta
from participa_geo.api.schemas.history import HistoryDTO
from participa_geo.api.schemas.scenarios import ScenarioDTO, ScenarioResultDTO
from participa_geo.reporting import maps as maps_module
from participa_geo.reporting.delivery import DeliveryContext, load_delivery_context
from participa_geo.reporting.geojson import export_scenario_geojson
from participa_geo.reporting.maps import (
    RISK_COLORS,
    MapRenderError,
    render_scenario_map,
)
from participa_geo.reporting.pdf import (
    generate_executive,
    generate_methodology,
    generate_scenario_pdf,
)
from participa_geo.reporting.tables import (
    export_history,
    export_model_evidence,
    export_scenario_csv,
)

REPO_ROOT = Path(__file__).resolve().parents[2]
MODEL_DIR = REPO_ROOT / "data" / "processed" / "model"
PROCESSED = REPO_ROOT / "data" / "processed"
REFERENCE = PROCESSED / "gis" / "reference-2025-12-15.geojson"

META = DatasetMeta(
    dataset_version_ids=[UUID("00000000-0000-5000-8000-000000000001")],
    reference_date="2025-12-15",
    model_run_id=UUID("5bc387ce-6646-5b49-9c4f-ad302f4c422a"),
    generated_at=datetime(2026, 10, 26, 12, 0, tzinfo=UTC),
    geographic_warning="Referencia cartográfica INE 2025-12-15.",
    filters={"territory_id": None, "year": 2024},
)


def dataset_meta(**overrides) -> DatasetMeta:
    values = META.model_dump()
    values.update(overrides)
    return DatasetMeta(**values)


def history_rows() -> list[HistoryDTO]:
    return [
        HistoryDTO(
            territory_id=UUID("00000000-0000-5000-8000-000000000002"),
            year=2024,
            name="ACATIC",
            municipality_code="001",
            ln=1000,
            sv=520,
            nv=430,
            ns=50,
            participation_pct=52.0,
            abstention_pct=43.0,
            unknown_share=0.05,
            ns_lower_pct=38.0,
            ns_upper_pct=48.0,
            eligible=True,
            identity_status="verified_catalogue",
            footprint_status="not_verified",
        ),
        HistoryDTO(
            territory_id=UUID("00000000-0000-5000-8000-000000000003"),
            year=2024,
            name="SIN DATO",
            municipality_code="002",
            ln=1000,
            sv=500,
            nv=400,
            ns=100,
            participation_pct=None,
            abstention_pct=None,
            unknown_share=None,
            ns_lower_pct=None,
            ns_upper_pct=None,
            eligible=False,
            identity_status="verified_catalogue",
            footprint_status="not_verified",
        ),
    ]


def _csv_rows(payload: bytes) -> tuple[list[str], list[dict[str, str]]]:
    text = payload.decode("utf-8-sig")
    lines = text.splitlines()
    data_start = next(
        index for index, line in enumerate(lines) if line and not line.startswith("#")
    )
    reader = csv.DictReader(io.StringIO("\n".join(lines[data_start:])))
    return reader.fieldnames or [], list(reader)


def test_csv_xlsx_preserve_nulls_and_versions():
    rows = history_rows()

    csv_payload = export_history(rows, "csv", META)
    text = csv_payload.decode("utf-8-sig")
    assert "# model_run_id=5bc387ce-6646-5b49-9c4f-ad302f4c422a" in text
    assert "# reference_date=2025-12-15" in text
    assert "00000000-0000-5000-8000-000000000001" in text

    fields, parsed = _csv_rows(csv_payload)
    assert "abstention_pct" in fields
    assert parsed[0]["abstention_pct"] == "43.0"
    assert parsed[1]["abstention_pct"] == ""
    assert parsed[1]["unknown_share"] == ""

    xlsx = export_history(rows, "xlsx", META)
    workbook = load_workbook(io.BytesIO(xlsx), data_only=False)
    assert "Datos" in workbook.sheetnames
    assert "Metadatos" in workbook.sheetnames
    assert "Diccionario" in workbook.sheetnames
    sheet = workbook["Datos"]
    headers = [cell.value for cell in sheet[1]]
    abstention_index = headers.index("abstention_pct") + 1
    assert sheet.cell(row=2, column=abstention_index).value == 43.0
    assert sheet.cell(row=3, column=abstention_index).value in (None, "")
    metadata = "\n".join(
        str(cell.value) for row in workbook["Metadatos"].iter_rows() for cell in row
    )
    assert "5bc387ce-6646-5b49-9c4f-ad302f4c422a" in metadata
    assert "2025-12-15" in metadata

    dictionary = workbook["Diccionario"]
    dictionary_headers = [cell.value for cell in dictionary[1]]
    assert dictionary_headers == [
        "campo",
        "etiqueta",
        "tipo/unidad",
        "descripcion",
    ]
    first = [cell.value for cell in dictionary[2]]
    assert first[0] == "territory_id"
    assert first[1] == "identificador de referencia (UUID)"
    assert first[2] == "texto"
    assert "Municipio de referencia" in first[3]


def test_spreadsheet_formula_text_is_escaped():
    rows = [
        HistoryDTO(
            territory_id=UUID("00000000-0000-5000-8000-000000000002"),
            year=2024,
            name="=SUM(A1:A2)",
            municipality_code="+001",
            ln=1000,
            sv=520,
            nv=430,
            ns=50,
            participation_pct=52.0,
            abstention_pct=43.0,
            unknown_share=0.05,
            ns_lower_pct=38.0,
            ns_upper_pct=48.0,
            eligible=True,
            identity_status="@verified",
            footprint_status="not_verified",
        )
    ]
    _, parsed = _csv_rows(export_history(rows, "csv", META))
    assert parsed[0]["name"].startswith("'=")
    assert parsed[0]["municipality_code"].startswith("'+")
    assert parsed[0]["identity_status"].startswith("'@")

    workbook = load_workbook(
        io.BytesIO(export_history(rows, "xlsx", META)), data_only=False
    )
    sheet = workbook["Datos"]
    headers = [cell.value for cell in sheet[1]]
    name_index = headers.index("name") + 1
    code_index = headers.index("municipality_code") + 1
    assert str(sheet.cell(row=2, column=name_index).value).startswith("'=")
    assert str(sheet.cell(row=2, column=code_index).value).startswith("'+")


def _real_scenario() -> ScenarioDTO:
    import pandas as pd

    def opt(value):
        return None if value is None or pd.isna(value) else float(value)

    frame = pd.read_parquet(MODEL_DIR / "predictions-2027.parquet")
    results: list[ScenarioResultDTO] = []
    selected = None
    for row in frame.itertuples(index=False):
        eligible = bool(row.eligible) and row.estimated_abstention_pct is not None
        baseline = float(row.estimated_abstention_pct) if eligible else None
        choose = eligible and selected is None
        if choose:
            selected = UUID(str(row.territory_id))
        hypothetical = max(0.0, baseline - 3.0) if choose and baseline is not None else baseline
        effective = (
            (baseline - hypothetical)
            if baseline is not None and hypothetical is not None
            else 0.0
        )
        results.append(
            ScenarioResultDTO(
                territory_id=UUID(str(row.territory_id)),
                municipality_code=str(row.municipality_code),
                municipality_name=str(row.municipality_name).upper(),
                baseline_abstention_pct=baseline,
                hypothetical_abstention_pct=hypothetical,
                baseline_error_lower_pct=opt(row.historical_error_lower_pct),
                baseline_error_upper_pct=opt(row.historical_error_upper_pct),
                hypothetical_error_lower_pct=(
                    max(0.0, float(row.historical_error_lower_pct) - 3.0)
                    if choose and not pd.isna(row.historical_error_lower_pct)
                    else opt(row.historical_error_lower_pct)
                ),
                hypothetical_error_upper_pct=(
                    min(100.0, float(row.historical_error_upper_pct) - 3.0)
                    if choose and not pd.isna(row.historical_error_upper_pct)
                    else opt(row.historical_error_upper_pct)
                ),
                requested_reduction_pp=3.0,
                effective_reduction_pp=effective,
                risk_a=str(row.risk),
                risk_b=str(row.risk),
                risk_a_unstable=bool(row.risk_unstable),
                risk_b_unstable=bool(row.risk_unstable),
                selected=choose,
            )
        )
    assert selected is not None
    return ScenarioDTO(
        scenario_id=UUID("9b2c5f1e-4d3a-4c6b-8e7f-1a2b3c4d5e6f"),
        model_run_id=UUID("5bc387ce-6646-5b49-9c4f-ad302f4c422a"),
        model_sha256="0" * 64,
        created_at=datetime(2026, 10, 26, 12, 0, tzinfo=UTC),
        target_year=2027,
        requested_reduction_pp=3.0,
        assumption_code="preventive_exploration",
        assumption_version="1.0",
        rationale_text="Supuesto de exploración preventiva revisado.",
        rules_version="scenario-rules-1.0.0",
        input_sha256="1" * 64,
        error_band_label=(
            "rango de error histórico de referencia trasladado condicionalmente; "
            "no es un intervalo causal"
        ),
        selected_territory_ids=[selected],
        results=results,
    )


def test_geojson_matches_scenario_and_reference():
    scenario = _real_scenario()
    payload = export_scenario_geojson(scenario, REFERENCE, META)
    body = json.loads(payload)
    assert body["type"] == "FeatureCollection"
    assert len(body["features"]) == 125
    assert body["metadata"]["reference_date"] == "2025-12-15"
    assert body["metadata"]["scenario_id"] == str(scenario.scenario_id)

    by_code = {result.municipality_code: result for result in scenario.results}
    seen = set()
    for feature in body["features"]:
        props = feature["properties"]
        code = props["municipality_code"]
        seen.add(code)
        result = by_code[code]
        assert props["baseline_abstention_pct"] == result.baseline_abstention_pct
        assert props["hypothetical_abstention_pct"] == result.hypothetical_abstention_pct
        if result.baseline_abstention_pct is None:
            assert props["baseline_abstention_pct"] is None
            assert props["hypothetical_abstention_pct"] is None
            assert props["effective_reduction_pp"] is None
        for ring in _all_coordinates(feature["geometry"]):
            x, y = ring
            assert -180.0 <= x <= 180.0
            assert -90.0 <= y <= 90.0
    assert seen == set(by_code)
    # The reference file is the same one used by the API map (125, cut 2025-12-15).
    reference = json.loads(REFERENCE.read_text(encoding="utf-8"))
    assert len(reference["features"]) == 125


def _all_coordinates(geometry: dict):
    if geometry["type"] == "Polygon":
        rings = geometry["coordinates"]
    elif geometry["type"] == "MultiPolygon":
        rings = [ring for polygon in geometry["coordinates"] for ring in polygon]
    else:  # pragma: no cover - the INE reference is polygon-based
        raise AssertionError(f"unexpected geometry {geometry['type']}")
    for ring in rings:
        for position in ring:
            yield position[0], position[1]


def test_scenario_csv_matches_api_values():
    scenario = _real_scenario()
    fields, rows = _csv_rows(export_scenario_csv(scenario, META))
    assert len(rows) == 125
    assert "baseline_abstention_pct" in fields
    assert "hypothetical_abstention_pct" in fields
    by_code = {row["municipality_code"]: row for row in rows}
    for result in scenario.results:
        row = by_code[result.municipality_code]
        assert row["scenario_id"] == str(scenario.scenario_id)
        if result.baseline_abstention_pct is None:
            assert row["baseline_abstention_pct"] == ""
            assert row["hypothetical_abstention_pct"] == ""
            assert row["effective_reduction_pp"] == ""
        else:
            assert float(row["baseline_abstention_pct"]) == pytest.approx(
                result.baseline_abstention_pct
            )


def test_scenario_selected_means_do_not_use_all_eligible():
    scenario = _real_scenario()
    context = load_delivery_context(scenario=scenario)
    assert context.scenario is not None
    selected = [row for row in scenario.results if row.selected]
    selected_a = [row.baseline_abstention_pct for row in selected]
    all_a = [
        row.baseline_abstention_pct
        for row in scenario.results
        if row.baseline_abstention_pct is not None
    ]
    assert len(all_a) > len(selected_a)
    expected_selected = sum(selected_a) / len(selected_a)
    expected_all = sum(all_a) / len(all_a)
    assert context.scenario.baseline_mean == pytest.approx(expected_selected)
    assert context.scenario.baseline_mean != pytest.approx(expected_all)
    assert context.scenario.selected_count == len(selected)


class _FakeClient:
    def __init__(self, *, fail: bool = False, body: bytes = b"\x89PNG\r\n\x1a\n"):
        self.fail = fail
        self.body = body
        self.calls: list[dict] = []

    def get_map(self, layer, scenario_id, year, bbox, width, height, **kwargs):
        self.calls.append(
            {
                "layer": layer,
                "scenario_id": scenario_id,
                "width": width,
                "height": height,
                "kwargs": kwargs,
            }
        )
        if self.fail:
            from participa_geo.api.gis_client import GeoServerError

            raise GeoServerError("GetMap failed")
        return self.body


def test_map_failure_is_not_placeholder_success():
    scenario = _real_scenario()
    with pytest.raises(MapRenderError):
        render_scenario_map(scenario, "png", client=_FakeClient(fail=True))
    with pytest.raises(MapRenderError):
        render_scenario_map(
            scenario,
            "png",
            client=_FakeClient(body=b"<?xml version='1.0'?><ServiceExceptionReport/>"),
        )


def test_map_requests_1200x900_with_shared_reference():
    scenario = _real_scenario()
    client = _FakeClient()
    with pytest.raises(MapRenderError):
        # A PNG signature only is not a decodable image; the failure is explicit.
        render_scenario_map(scenario, "png", client=client)
    assert len(client.calls) == 2
    for call in client.calls:
        assert call["width"] == maps_module.MAP_WIDTH
        assert call["height"] == maps_module.MAP_HEIGHT


def test_map_rejects_wrong_png_dimensions():
    scenario = _real_scenario()
    buffer = io.BytesIO()
    Image.new("RGB", (600, 450), "#FFFFFF").save(buffer, format="PNG")
    with pytest.raises(MapRenderError, match="1200×900"):
        render_scenario_map(
            scenario, "png", client=_FakeClient(body=buffer.getvalue())
        )


def test_ab_sld_legend_and_insufficient_style_are_identical():
    risk = (REPO_ROOT / "gis" / "styles" / "risk.sld").read_text(encoding="utf-8")
    risk_a = (REPO_ROOT / "gis" / "styles" / "risk_a.sld").read_text(
        encoding="utf-8"
    )

    def rule_symbols(text: str, name: str) -> dict[str, str]:
        match = re.search(rf"<Name>{name}</Name>(.*?)</Rule>", text, re.S)
        assert match, f"rule {name} not found"
        block = match.group(1)
        return {
            "fill": re.search(r'name="fill">([^<]+)', block).group(1).upper(),  # type: ignore[union-attr]
            "opacity": re.search(r'name="fill-opacity">([^<]+)', block).group(1),  # type: ignore[union-attr]
            "stroke": re.search(r'name="stroke">([^<]+)', block).group(1).upper(),  # type: ignore[union-attr]
            "stroke-width": re.search(r'name="stroke-width">([^<]+)', block).group(1),  # type: ignore[union-attr]
        }

    for name in ("low", "medium", "high", "insufficient"):
        assert rule_symbols(risk_a, name) == rule_symbols(risk, name), name
        assert rule_symbols(risk_a, name)["fill"] == RISK_COLORS[name].upper()


def test_methodology_page_count(delivery_context: DeliveryContext):
    report = generate_methodology(delivery_context)
    pages = len(PdfReader(io.BytesIO(report)).pages)
    assert 10 <= pages <= 15


def test_executive_page_count(delivery_context: DeliveryContext):
    report = generate_executive(delivery_context)
    pages = len(PdfReader(io.BytesIO(report)).pages)
    assert pages <= 10


def test_pdf_text_contains_real_results_and_assumptions(delivery_context: DeliveryContext):
    text = "\n".join(
        page.extract_text() or ""
        for page in PdfReader(io.BytesIO(generate_methodology(delivery_context))).pages
    )
    assert "M0" in text
    assert "4.93" in text
    assert "8.90" in text
    assert "3.43" in text
    assert "8.56" in text
    assert "85" in text and "40" in text
    assert "70" in text
    assert "no implica efecto causal" in text.lower()
    assert "2025-12-15" in text
    assert "fuentes" in text.lower()
    assert "Explicación aditiva del candidato M1" in text
    assert "model_evidence_csv" in text

    executive = "\n".join(
        page.extract_text() or ""
        for page in PdfReader(io.BytesIO(generate_executive(delivery_context))).pages
    )
    assert "85" in executive and "40" in executive
    assert "4.93" in executive
    assert "no causal" in executive.lower() or "no implica" in executive.lower()


def test_pdf_numbers_are_derived_from_context(delivery_context: DeliveryContext):
    mutated = replace(
        delivery_context,
        cohort=replace(
            delivery_context.cohort,
            total=200,
            eligible=101,
            insufficient=99,
            unstable=66,
        ),
    )
    methodology = "\n".join(
        page.extract_text() or ""
        for page in PdfReader(io.BytesIO(generate_methodology(mutated))).pages
    )
    assert "elegibles de 200" in methodology
    assert "101" in methodology
    assert "99" in methodology
    assert "66" in methodology
    assert "1200 observaciones municipales" in methodology
    assert "70 de los 85" not in methodology
    assert "750 observaciones" not in methodology

    executive = "\n".join(
        page.extract_text() or ""
        for page in PdfReader(io.BytesIO(generate_executive(mutated))).pages
    )
    assert "101" in executive
    assert "99" in executive
    assert "66" in executive
    assert "1200" in executive
    assert "750" not in executive


def test_scenario_effective_reduction_is_reported_as_range():
    scenario = _real_scenario()
    clipped_results = [
        row.model_copy(
            update={
                "baseline_abstention_pct": 1.0,
                "hypothetical_abstention_pct": 0.0,
                "effective_reduction_pp": 1.0,
            }
        )
        if row.selected
        else row
        for row in scenario.results
    ]
    clipped = scenario.model_copy(update={"results": clipped_results})
    context = load_delivery_context(scenario=clipped)
    assert context.scenario is not None
    assert context.scenario.effective_reduction_min_pp == pytest.approx(1.0)
    assert context.scenario.effective_reduction_max_pp == pytest.approx(1.0)

    text = "\n".join(
        page.extract_text() or ""
        for page in PdfReader(io.BytesIO(generate_methodology(context))).pages
    )
    assert "3.0 pp / 1.0 pp (rango seleccionado)" in text
    scenario_text = "\n".join(
        page.extract_text() or ""
        for page in PdfReader(
            io.BytesIO(generate_scenario_pdf(clipped, META, context=context))
        ).pages
    )
    assert "3.0 pp / 1.0 pp" in scenario_text


def test_pdf_exclusion_reasons_are_derived_per_cut(delivery_context: DeliveryContext):
    from participa_geo.reporting.delivery import ExclusionFact

    mutated = replace(
        delivery_context,
        cohort=replace(
            delivery_context.cohort,
            excluded_by_year={**delivery_context.cohort.excluded_by_year, 2027: 3},
            exclusion_details={
                **delivery_context.exclusion_details,
                2027: (
                    ExclusionFact("001", ("high_unknown_share", "missing_lag")),
                    ExclusionFact("002", ("ambiguous_identity",)),
                    ExclusionFact("003", ("known_administrative_change",)),
                ),
            },
        ),
    )
    text = "\n".join(
        page.extract_text() or ""
        for page in PdfReader(io.BytesIO(generate_methodology(mutated))).pages
    )
    assert "high_unknown_share (1)" in text
    assert "missing_lag (1)" in text
    assert "ambiguous_identity (1)" in text
    assert "known_administrative_change (1)" in text


def test_pdf_labels_2027_ln_coverage_with_2024_base(
    delivery_context: DeliveryContext,
):
    assert delivery_context.cohort.ln_coverage_base_year == 2024
    text = "\n".join(
        page.extract_text() or ""
        for page in PdfReader(io.BytesIO(generate_methodology(delivery_context))).pages
    )
    assert "LN observada de 2024" in text
    assert "antecedente" in text
    assert "proyección de la LN 2027" in text


def test_pdf_reports_empirical_2024_coverage_without_changing_amplitude(
    delivery_context: DeliveryContext,
):
    covered = replace(
        delivery_context,
        model=replace(
            delivery_context.model,
            empirical_coverage_2024={
                "n": 75,
                "within": 40,
                "share": 40 / 75,
                "amplitude_pp": delivery_context.model.band_pp,
                "minimum_n": 30,
            },
        ),
    )
    text = "\n".join(
        page.extract_text() or ""
        for page in PdfReader(io.BytesIO(generate_methodology(covered))).pages
    )
    assert "40 de 75" in text
    assert "53.3%" in text
    assert f"{delivery_context.model.band_pp:.2f}" in text

    too_small = replace(
        delivery_context,
        model=replace(delivery_context.model, empirical_coverage_2024=None),
    )
    small_text = "\n".join(
        page.extract_text() or ""
        for page in PdfReader(io.BytesIO(generate_methodology(too_small))).pages
    )
    assert "menos de 30" in small_text


def test_model_evidence_csv_xlsx_contains_features_exclusions_and_contributions(
    tmp_path: Path,
):
    model_dir = tmp_path / "model"
    model_dir.mkdir()
    (model_dir / "model.json").write_text(
        json.dumps(
            {
                "family": "M0",
                "target_year": 2027,
                "model_run_id": "00000000-0000-5000-8000-0000000000aa",
                "feature_order": [
                    "same_type_abstention_pct",
                    "is_midterm",
                ],
                "features_2027": [
                    {
                        "municipality_code": "001",
                        "municipality_name": "ACATIC",
                        "features": {
                            "same_type_abstention_pct": 47.5,
                            "is_midterm": 1.0,
                        },
                        "origins": {
                            "same_type_abstention_pct": 2021,
                            "is_midterm": None,
                        },
                    }
                ],
                "explanations": [
                    {
                        "municipality_code": "001",
                        "intercept": 0.0,
                        "contributions": {"same_type_abstention_pct": 47.5},
                        "pre_clip_sum": 47.5,
                    }
                ],
                "candidate_m1_explanations": [
                    {
                        "municipality_code": "001",
                        "intercept": 36.3,
                        "contributions": {"same_type_abstention_pct": 5.2},
                        "pre_clip_sum": 41.5,
                    }
                ],
            }
        ),
        encoding="utf-8",
    )
    (model_dir / "evaluation.json").write_text(
        json.dumps(
            {
                "exclusions": {
                    "2027": [
                        {
                            "municipality_code": "001",
                            "reasons": ["high_unknown_share", "missing_lag"],
                        }
                    ],
                    "2024": [
                        {
                            "municipality_code": "002",
                            "reasons": ["ambiguous_identity"],
                        }
                    ],
                }
            }
        ),
        encoding="utf-8",
    )

    csv_payload = export_model_evidence("csv", model_dir, META)
    fields, rows = _csv_rows(csv_payload)
    assert fields == [
        "record_type",
        "cut_year",
        "municipality_code",
        "municipality_name",
        "field",
        "value",
        "unit",
        "source_year",
        "reason",
        "model_run_id",
    ]
    kinds = {row["record_type"] for row in rows}
    assert kinds == {
        "feature",
        "exclusion",
        "contribution_selected",
        "contribution_candidate_m1",
    }
    by_key = {
        (
            row["record_type"],
            row["cut_year"],
            row["municipality_code"],
            row["field"],
        ): row
        for row in rows
    }
    assert (
        by_key[("feature", "2027", "001", "same_type_abstention_pct")]["value"]
        == "47.5"
    )
    exclusion = by_key[("exclusion", "2027", "001", "")]
    assert "high_unknown_share" in exclusion["reason"]
    assert "missing_lag" in exclusion["reason"]
    assert (
        by_key[("contribution_candidate_m1", "2027", "001", "same_type_abstention_pct")][
            "value"
        ]
        == "5.2"
    )

    workbook = load_workbook(
        io.BytesIO(export_model_evidence("xlsx", model_dir, META)),
        data_only=False,
    )
    assert {"Evidencia", "Metadatos", "Diccionario"} <= set(workbook.sheetnames)
    evidence = workbook["Evidencia"]
    headers = [cell.value for cell in evidence[1]]
    assert headers[0] == "record_type"
    assert evidence.max_row == len(rows) + 1


def test_delivery_context_is_built_from_artifacts():
    context = load_delivery_context(MODEL_DIR, PROCESSED)
    assert context.model.family == "M0"
    assert context.model.mae_2021 == pytest.approx(4.93363617513041)
    assert context.model.m1_mae_2021 == pytest.approx(8.89982962915524)
    assert context.model.band_pp == pytest.approx(8.561771770259021)
    assert context.cohort.total == 125
    assert context.cohort.eligible == 85
    assert context.cohort.insufficient == 40
    assert context.cohort.unstable == 70
    assert context.reference_date == "2025-12-15"
    assert len(context.sources) == 6
