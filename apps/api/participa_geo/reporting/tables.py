"""CSV/XLSX exports for municipal history, scenarios and dictionaries (T9).

Nulls stay empty (never zero), every export carries dataset/model/version
metadata, and text cells that a spreadsheet could interpret as a formula are
prefixed with an apostrophe in both CSV and XLSX. The XLSX workbook always has a
data sheet, a metadata sheet and a dictionary sheet.
"""

# Long Spanish dictionary labels/descriptions are intentional editorial text.
# ruff: noqa: E501

from __future__ import annotations

import csv
import io
import json
from collections.abc import Iterable, Sequence
from datetime import datetime
from pathlib import Path
from typing import Any

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from participa_geo.api.schemas.common import DatasetMeta
from participa_geo.api.schemas.history import HistoryDTO
from participa_geo.api.schemas.reports import REPORT_RULES_VERSION
from participa_geo.api.schemas.scenarios import ScenarioDTO

_FORMULA_PREFIXES = ("=", "+", "-", "@", "\t", "\r")


def _escape_spreadsheet(value: Any) -> Any:
    """Neutralise a leading spreadsheet formula prefix in string values."""
    if not isinstance(value, str):
        return value
    stripped = value.lstrip()
    if stripped[:1] in _FORMULA_PREFIXES or value[:1] in ("\t", "\r"):
        return "'" + value
    return value


def _cell(value: Any) -> str:
    value = _escape_spreadsheet(value)
    if value is None:
        return ""
    if isinstance(value, bool):
        return "true" if value else "false"
    return str(value)


def _excel_value(value: Any) -> Any:
    """Convert UUID/datetime strings for openpyxl and escape formulas."""
    if value is None or isinstance(value, (int, float, bool)):
        return value
    if isinstance(value, datetime):
        return value.isoformat()
    return _escape_spreadsheet(str(value))


def _json_value(value: Any) -> str:
    if isinstance(value, datetime):
        return value.isoformat()
    return json.dumps(value, ensure_ascii=False, sort_keys=True, default=str)


HISTORY_COLUMNS: tuple[tuple[str, str, str, str], ...] = (
    (
        "territory_id",
        "identificador de referencia (UUID)",
        "texto",
        "Municipio de referencia INE 2025; vacío en agregados estatales.",
    ),
    ("year", "año", "entero", "Año electoral CCPC (2009, 2012, 2015, 2018, 2021, 2024)."),
    ("name", "municipio", "texto", "Nombre del municipio o nulo en el agregado estatal."),
    ("municipality_code", "código municipal", "texto", "Código INE de tres dígitos."),
    ("ln", "lista nominal", "personas", "Lista nominal agregada de la fuente CCPC."),
    ("sv", "votación", "personas", "Sufragios válidos agregados."),
    ("nv", "no votó", "personas", "Personas en lista nominal que no votaron."),
    ("ns", "no especificado", "personas", "Conteo no especificado de la fuente."),
    (
        "participation_pct",
        "participación",
        "porcentaje",
        "100 × SV / (SV + NV); nulo si el denominador es cero.",
    ),
    (
        "abstention_pct",
        "abstención",
        "porcentaje",
        "100 × NV / (SV + NV); nulo si el denominador es cero.",
    ),
    (
        "unknown_share",
        "proporción no especificada",
        "proporción 0–1",
        "NS / LN; nulo si LN es cero.",
    ),
    ("ns_lower_pct", "cota inferior NS", "porcentaje", "100 × NV / LN; supuesto de cota inferior."),
    (
        "ns_upper_pct",
        "cota superior NS",
        "porcentaje",
        "100 × (NV + NS) / LN; supuesto de cota superior.",
    ),
    (
        "eligible",
        "elegible",
        "booleano",
        "Decisión de calidad: alta proporción de no especificados excluye.",
    ),
    (
        "identity_status",
        "estado de identidad",
        "catálogo",
        "Coincidencia de catálogo entre fuente y referencia.",
    ),
    (
        "footprint_status",
        "estado de huella",
        "catálogo",
        "Igualdad histórica de límites: no verificada, nunca inferida.",
    ),
    ("level", "nivel geográfico", "catálogo", "municipality o federal_district."),
    (
        "territory_code",
        "clave territorial",
        "texto",
        "Clave INE del territorio; distrito asociado a plan.",
    ),
    (
        "district_plan",
        "plan distrital",
        "catálogo",
        "DERFE2023 para la observación distrital 2024.",
    ),
    (
        "coverage_share",
        "cobertura distrital",
        "proporción 0–1",
        "Cobertura observada de LN 2024; no representa cobertura predictiva.",
    ),
)

HISTORY_FIELDS: tuple[str, ...] = tuple(column[0] for column in HISTORY_COLUMNS)

SCENARIO_COLUMNS: tuple[tuple[str, str, str, str], ...] = (
    (
        "scenario_id",
        "identificador del escenario",
        "UUID",
        "Escenario persistido; no se recalcula al exportar.",
    ),
    ("territory_id", "identificador de referencia", "UUID", "Municipio de referencia INE 2025."),
    ("municipality_code", "código municipal", "texto", "Código INE de tres dígitos."),
    ("municipality_name", "municipio", "texto", "Nombre publicado."),
    (
        "selected",
        "seleccionado",
        "booleano",
        "Verdadero sólo si el municipio recibió la reducción hipotética.",
    ),
    (
        "baseline_abstention_pct",
        "A · abstención estimada",
        "porcentaje",
        "Predicción publicada 2027; nulo si es insuficiente.",
    ),
    (
        "hypothetical_abstention_pct",
        "B · abstención hipotética",
        "porcentaje",
        "max(0, A − reducción) en seleccionados; nulo si es insuficiente.",
    ),
    (
        "baseline_error_lower_pct",
        "A · cota inferior",
        "porcentaje",
        "Banda de error histórico de referencia; no es intervalo causal.",
    ),
    (
        "baseline_error_upper_pct",
        "A · cota superior",
        "porcentaje",
        "Banda de error histórico de referencia; no es intervalo causal.",
    ),
    (
        "hypothetical_error_lower_pct",
        "B · cota inferior",
        "porcentaje",
        "Banda A trasladada condicionalmente y recortada a 0.",
    ),
    (
        "hypothetical_error_upper_pct",
        "B · cota superior",
        "porcentaje",
        "Banda A trasladada condicionalmente y recortada a 100.",
    ),
    (
        "requested_reduction_pp",
        "reducción solicitada",
        "puntos porcentuales",
        "Valor original de la solicitud (0–10).",
    ),
    (
        "effective_reduction_pp",
        "reducción efectiva",
        "puntos porcentuales",
        "A − B; puede ser menor si B se recorta a cero.",
    ),
    (
        "risk_a",
        "riesgo A",
        "categoría",
        "Semáforo analítico: low/medium/high/insufficient; no es clasificación oficial.",
    ),
    ("risk_b", "riesgo B", "categoría", "Mismos cortes que A; no es clasificación oficial."),
    ("risk_a_unstable", "A inestable", "booleano", "La banda cruza un umbral de riesgo."),
    (
        "risk_b_unstable",
        "B inestable",
        "booleano",
        "La banda trasladada cruza un umbral de riesgo.",
    ),
    ("assumption_code", "supuesto", "catálogo", "preventive_exploration o model_sensitivity."),
    (
        "assumption_version",
        "versión del supuesto",
        "texto",
        "Versión controlada del texto de supuesto.",
    ),
    ("rules_version", "versión de reglas", "texto", "Versión de las reglas de escenario."),
    ("target_year", "año objetivo", "entero", "Fijo en 2027."),
    (
        "input_sha256",
        "hash de entrada",
        "sha256",
        "Huella de las entradas del escenario persistido.",
    ),
    ("model_run_id", "ejecución de modelo", "UUID", "Modelo publicado utilizado."),
    ("model_sha256", "hash del modelo", "sha256", "Huella lógica del modelo."),
)

SCENARIO_FIELDS: tuple[str, ...] = tuple(column[0] for column in SCENARIO_COLUMNS)

MODEL_EVIDENCE_COLUMNS: tuple[tuple[str, str, str, str], ...] = (
    (
        "record_type",
        "tipo de evidencia",
        "catálogo",
        "feature, exclusion, contribution_selected o contribution_candidate_m1.",
    ),
    ("cut_year", "corte", "entero/año", "Año del corte; 2027 para features y contribuciones."),
    ("state_code", "clave de entidad", "texto", "Clave INE de entidad de dos dígitos."),
    ("state_name", "entidad", "texto", "Nombre oficial de la entidad."),
    ("municipality_code", "código municipal", "texto", "Código INE de tres dígitos."),
    ("municipality_name", "municipio", "texto", "Nombre publicado."),
    ("field", "campo", "texto", "Variable o motivo de exclusión."),
    ("value", "valor", "según campo", "Valor de la feature o contribución aditiva."),
    ("unit", "unidad", "texto", "porcentaje, proporción 0–1, indicador o puntos porcentuales."),
    (
        "source_year",
        "año de origen",
        "entero o vacío",
        "Año de la observación que origina la feature; vacío si es estructural.",
    ),
    (
        "reason",
        "motivo de exclusión",
        "texto",
        "Motivos estables separados por punto y coma; vacío en features/contribuciones.",
    ),
    ("model_run_id", "ejecución de modelo", "UUID", "Modelo publicado que origina la evidencia."),
)

MODEL_EVIDENCE_FIELDS: tuple[str, ...] = tuple(column[0] for column in MODEL_EVIDENCE_COLUMNS)
MAX_MODEL_EVIDENCE_ROWS = 20_000


class ModelEvidenceLimitError(RuntimeError):
    """Model evidence exceeds the documented maximum; no rows were truncated."""

_FEATURE_UNITS: dict[str, str] = {
    "same_type_abstention_pct": "porcentaje",
    "previous_abstention_pct": "porcentaje",
    "previous_young_share": "proporción 0–1",
    "previous_senior_share": "proporción 0–1",
    "previous_rural_share": "proporción 0–1",
    "is_midterm": "indicador 0/1",
}

REPORT_DICTIONARY: tuple[tuple[str, str, str, str], ...] = (
    (
        "report_artifact",
        "report_id",
        "UUID",
        "Identificador determinista del artefacto; la URL de descarga no expone rutas.",
    ),
    (
        "report_artifact",
        "sha256",
        "sha256",
        "Hash del contenido descargado; permite verificar integridad.",
    ),
    ("report_artifact", "status", "catálogo", "ready o failed; no existe estado pending ni cola."),
    (
        "report_artifact",
        "cache_key",
        "sha256 (interno)",
        "Filtros, versiones, modelo, escenario y supuestos; una caché sólo se reutiliza si coincide.",
    ),
    (
        "scenario",
        "error_band_label",
        "texto",
        "Banda de error histórico de referencia; no es intervalo causal.",
    ),
)


def _metadata_lines(meta: DatasetMeta, sources: dict | None) -> list[str]:
    lines = [
        "# PARTICIPA-GEO · exportación nacional agregada",
        f"# kind={meta.filters.get('report_kind', '')}",
        f"# generated_at={meta.generated_at.isoformat()}",
        f"# reference_date={meta.reference_date}",
        f"# model_run_id={meta.model_run_id or ''}",
        f"# dataset_version_ids={_json_value([str(item) for item in meta.dataset_version_ids])}",
        f"# geographic_warning={meta.geographic_warning}",
        f"# filters={_json_value(meta.filters)}",
        f"# report_rules_version={REPORT_RULES_VERSION}",
    ]
    if sources:
        lines.append(f"# sources={_json_value(sources)}")
    lines.append(
        "# uso=agregados municipales; prohibida la inferencia individual; "
        "el escenario es hipotético y no implica efecto causal"
    )
    return lines


def _resolve_kind(kind: str) -> str:
    value = kind.lower()
    if value not in {"csv", "xlsx"}:
        raise ValueError("kind debe ser 'csv' o 'xlsx'")
    return value


def export_history(
    rows: Sequence[HistoryDTO],
    kind: str,
    meta: DatasetMeta,
    *,
    sources: dict | None = None,
) -> bytes:
    """Export municipal/state history as CSV or a three-sheet XLSX."""
    fmt = _resolve_kind(kind)
    if fmt == "csv":
        buffer = io.StringIO()
        for line in _metadata_lines(meta, sources):
            buffer.write(line + "\n")
        writer = csv.writer(buffer, lineterminator="\n")
        writer.writerow(HISTORY_FIELDS)
        for row in rows:
            payload = row.model_dump(mode="python")
            writer.writerow([_cell(payload.get(field)) for field in HISTORY_FIELDS])
        return buffer.getvalue().encode("utf-8-sig")

    workbook = Workbook()
    data = workbook.active
    data.title = "Datos"
    headers = list(HISTORY_FIELDS)
    data.append(headers)
    for row in rows:
        payload = row.model_dump(mode="python")
        data.append([_excel_value(payload.get(field)) for field in HISTORY_FIELDS])
    _style_data_sheet(data, len(headers))

    metadata = workbook.create_sheet("Metadatos")
    _write_metadata_sheet(metadata, meta, sources)
    dictionary = workbook.create_sheet("Diccionario")
    _write_dictionary_sheet(dictionary, HISTORY_COLUMNS)

    stream = io.BytesIO()
    workbook.save(stream)
    return stream.getvalue()


def export_scenario_csv(
    scenario: ScenarioDTO,
    meta: DatasetMeta,
    *,
    sources: dict | None = None,
) -> bytes:
    """Export every persisted scenario A/B row, matching API values exactly."""
    buffer = io.StringIO()
    metadata = meta.model_copy(
        update={"filters": {**meta.filters, "scenario_id": str(scenario.scenario_id)}}
    )
    for line in _metadata_lines(metadata, sources):
        buffer.write(line + "\n")
    buffer.write(f"# assumption_code={scenario.assumption_code}\n")
    buffer.write(f"# assumption_version={scenario.assumption_version}\n")
    buffer.write(f"# rules_version={scenario.rules_version}\n")
    buffer.write(f"# rationale_text={scenario.rationale_text}\n")
    buffer.write(f"# input_sha256={scenario.input_sha256}\n")
    buffer.write(f"# model_sha256={scenario.model_sha256}\n")
    buffer.write(f"# error_band_label={scenario.error_band_label}\n")
    writer = csv.writer(buffer, lineterminator="\n")
    writer.writerow(SCENARIO_FIELDS)
    for result in scenario.results:
        writer.writerow(
            [
                _cell(str(scenario.scenario_id)),
                _cell(str(result.territory_id)),
                _cell(result.municipality_code),
                _cell(result.municipality_name),
                _cell(result.selected),
                _cell(result.baseline_abstention_pct),
                _cell(result.hypothetical_abstention_pct),
                _cell(result.baseline_error_lower_pct),
                _cell(result.baseline_error_upper_pct),
                _cell(result.hypothetical_error_lower_pct),
                _cell(result.hypothetical_error_upper_pct),
                _cell(result.requested_reduction_pp),
                _cell(
                    result.effective_reduction_pp
                    if result.baseline_abstention_pct is not None
                    else None
                ),
                _cell(result.risk_a),
                _cell(result.risk_b),
                _cell(result.risk_a_unstable),
                _cell(result.risk_b_unstable),
                _cell(scenario.assumption_code),
                _cell(scenario.assumption_version),
                _cell(scenario.rules_version),
                _cell(scenario.target_year),
                _cell(scenario.input_sha256),
                _cell(str(scenario.model_run_id)),
                _cell(scenario.model_sha256),
            ]
        )
    return buffer.getvalue().encode("utf-8-sig")


def _load_model_evidence(model_dir: Path | str) -> tuple[dict, dict]:
    model_dir = Path(model_dir)
    model = json.loads((model_dir / "model.json").read_text(encoding="utf-8"))
    evaluation = json.loads((model_dir / "evaluation.json").read_text(encoding="utf-8"))
    return model, evaluation


def _model_evidence_rows(
    model: dict, evaluation: dict, model_dir: Path | None = None
) -> list[dict[str, Any]]:
    model_run_id = str(model.get("model_run_id", ""))
    features = model.get("features_2027") or []
    if not features and model_dir and model.get("level") == "municipality":
        return _national_model_evidence_rows(model, model_dir)
    names = {
        str(item.get("municipality_code", "")): str(item.get("municipality_name", ""))
        for item in features
    }
    rows: list[dict[str, Any]] = []
    for item in features:
        code = str(item.get("municipality_code", ""))
        origins = item.get("origins") or {}
        for field, value in (item.get("features") or {}).items():
            rows.append(
                {
                    "record_type": "feature",
                    "cut_year": int(model.get("target_year", 2027)),
                    "state_code": item.get("state_code", ""),
                    "state_name": item.get("state_name", ""),
                    "municipality_code": code,
                    "municipality_name": item.get("municipality_name", ""),
                    "field": field,
                    "value": value,
                    "unit": _FEATURE_UNITS.get(field, "escala transformada"),
                    "source_year": origins.get(field),
                    "reason": "",
                    "model_run_id": model_run_id,
                }
            )
    for year, exclusions in sorted(
        (evaluation.get("exclusions") or {}).items(),
        key=lambda item: int(item[0]),
    ):
        for exclusion in exclusions:
            code = str(exclusion.get("municipality_code", ""))
            rows.append(
                {
                    "record_type": "exclusion",
                    "cut_year": int(year),
                    "state_code": exclusion.get("state_code", ""),
                    "state_name": exclusion.get("state_name", ""),
                    "municipality_code": code,
                    "municipality_name": names.get(code, ""),
                    "field": "",
                    "value": "",
                    "unit": "",
                    "source_year": "",
                    "reason": ";".join(exclusion.get("reasons", [])),
                    "model_run_id": model_run_id,
                }
            )
    for item in model.get("explanations") or []:
        for field, value in (item.get("contributions") or {}).items():
            rows.append(
                {
                    "record_type": "contribution_selected",
                    "cut_year": int(model.get("target_year", 2027)),
                    "state_code": item.get("state_code", ""),
                    "state_name": item.get("state_name", ""),
                    "municipality_code": item.get("municipality_code", ""),
                    "municipality_name": names.get(item.get("municipality_code", ""), ""),
                    "field": field,
                    "value": value,
                    "unit": "puntos porcentuales (escala transformada)",
                    "source_year": "",
                    "reason": "",
                    "model_run_id": model_run_id,
                }
            )
    for item in model.get("candidate_m1_explanations") or []:
        for field, value in (item.get("contributions") or {}).items():
            rows.append(
                {
                    "record_type": "contribution_candidate_m1",
                    "cut_year": int(model.get("target_year", 2027)),
                    "state_code": item.get("state_code", ""),
                    "state_name": item.get("state_name", ""),
                    "municipality_code": item.get("municipality_code", ""),
                    "municipality_name": names.get(item.get("municipality_code", ""), ""),
                    "field": field,
                    "value": value,
                    "unit": "puntos porcentuales (escala transformada)",
                    "source_year": "",
                    "reason": "",
                    "model_run_id": model_run_id,
                }
            )
    return rows


def _national_model_evidence_rows(model: dict, model_dir: Path) -> list[dict[str, Any]]:
    """Derive evidence from the immutable national observation/prediction artifacts."""
    import pandas as pd

    root = model_dir.parent
    observations = pd.read_csv(root / "municipal-observations.csv")
    predictions = pd.read_parquet(model_dir / "predictions-2027.parquet")
    model_run_id = str(model.get("model_run_id", ""))
    years = observations.loc[observations["year"].isin([2021, 2024])].copy()
    years["state_code"] = years["state_code"].astype(str).str.zfill(2)
    years["municipality_code"] = years["municipality_code"].astype(str).str.zfill(3)
    predictions["state_code"] = predictions["state_code"].astype(str).str.zfill(2)
    predictions["municipality_code"] = predictions["municipality_code"].astype(str).str.zfill(3)
    index = years.set_index(["state_code", "municipality_code", "year"])
    names = {
        (str(row.state_code).zfill(2), str(row.municipality_code).zfill(3)): row.municipality_name
        for row in observations.itertuples()
    }
    state_names = {
        str(row.state_code).zfill(2): str(row.state_name)
        for row in observations.itertuples()
    }
    rows: list[dict[str, Any]] = []
    for prediction in predictions.itertuples():
        key = (prediction.state_code, prediction.municipality_code)
        if not bool(prediction.eligible):
            continue
        try:
            same = index.loc[(*key, 2021)]
            previous = index.loc[(*key, 2024)]
        except KeyError:
            continue
        denominator_same = float(same.sv + same.nv)
        denominator_previous = float(previous.sv + previous.nv)
        if not denominator_same or not denominator_previous:
            continue
        feature_values = {
            "same_type_abstention_pct": float(same.nv) * 100 / denominator_same,
            "previous_abstention_pct": float(previous.nv) * 100 / denominator_previous,
            "previous_young_share": float(previous.young_share),
            "previous_senior_share": float(previous.senior_share),
            "previous_rural_share": float(previous.rural_share),
            "is_midterm": 1.0,
        }
        for field, value in feature_values.items():
            rows.append(
                {
                    "record_type": "feature",
                    "cut_year": 2027,
                    "state_code": prediction.state_code,
                    "state_name": state_names.get(prediction.state_code, ""),
                    "municipality_code": prediction.municipality_code,
                    "municipality_name": prediction.municipality_name,
                    "field": field,
                    "value": value,
                    "unit": _FEATURE_UNITS.get(field, "indicador 0/1"),
                    "source_year": 2021
                    if field == "same_type_abstention_pct"
                    else 2024
                    if field != "is_midterm"
                    else "",
                    "reason": "",
                    "model_run_id": model_run_id,
                }
            )
        rows.append(
            {
                "record_type": "contribution_selected",
                "cut_year": 2027,
                "state_code": prediction.state_code,
                "state_name": state_names.get(prediction.state_code, ""),
                "municipality_code": prediction.municipality_code,
                "municipality_name": prediction.municipality_name,
                "field": "same_type_abstention_pct",
                "value": feature_values["same_type_abstention_pct"],
                "unit": "puntos porcentuales",
                "source_year": 2021,
                "reason": "",
                "model_run_id": model_run_id,
            }
        )
    exclusions_path = model_dir / "exclusions-by-cut.csv"
    if exclusions_path.is_file():
        for exclusion in pd.read_csv(exclusions_path).itertuples():
            key = (str(exclusion.state_code).zfill(2), str(exclusion.municipality_code).zfill(3))
            rows.append(
                {
                    "record_type": "exclusion",
                    "cut_year": int(exclusion.year),
                    "state_code": key[0],
                    "state_name": state_names.get(key[0], ""),
                    "municipality_code": key[1],
                    "municipality_name": names.get(key, ""),
                    "field": "",
                    "value": "",
                    "unit": "",
                    "source_year": "",
                    "reason": str(exclusion.reasons),
                    "model_run_id": model_run_id,
                }
            )
    return rows


def export_model_evidence(
    kind: str,
    model_dir: Path | str,
    meta: DatasetMeta,
    *,
    sources: dict | None = None,
) -> bytes:
    """User-facing CSV/XLSX with features, per-cut exclusions and contributions."""
    fmt = _resolve_kind(kind)
    model, evaluation = _load_model_evidence(model_dir)
    rows = _model_evidence_rows(model, evaluation, Path(model_dir))
    if len(rows) > MAX_MODEL_EVIDENCE_ROWS:
        raise ModelEvidenceLimitError(
            f"La evidencia del modelo supera el límite de {MAX_MODEL_EVIDENCE_ROWS:,} filas; "
            "no se generó una exportación truncada."
        )
    national_model = model.get("level") == "municipality" and not model.get("features_2027")
    evidence_meta = meta.model_copy(
        update={
            "filters": {
                **meta.filters,
                "report_kind": f"model_evidence_{fmt}",
                "model_family": model.get("family"),
                **(
                    {
                        "candidate_m1_detail_omitted": (
                            "Per-municipality candidate M1 contributions are omitted to keep "
                            "the national evidence export bounded; model coefficients remain "
                            "in the technical model artifact."
                        )
                    }
                    if national_model
                    else {}
                ),
            }
        }
    )
    if fmt == "csv":
        buffer = io.StringIO()
        for line in _metadata_lines(evidence_meta, sources):
            buffer.write(line + "\n")
        writer = csv.writer(buffer, lineterminator="\n")
        writer.writerow(MODEL_EVIDENCE_FIELDS)
        for row in rows:
            writer.writerow([_cell(row.get(field)) for field in MODEL_EVIDENCE_FIELDS])
        return buffer.getvalue().encode("utf-8-sig")

    workbook = Workbook()
    data = workbook.active
    data.title = "Evidencia"
    data.append(list(MODEL_EVIDENCE_FIELDS))
    for row in rows:
        data.append([_excel_value(row.get(field)) for field in MODEL_EVIDENCE_FIELDS])
    _style_data_sheet(data, len(MODEL_EVIDENCE_FIELDS))
    metadata = workbook.create_sheet("Metadatos")
    _write_metadata_sheet(metadata, evidence_meta, sources)
    dictionary = workbook.create_sheet("Diccionario")
    _write_dictionary_sheet(dictionary, MODEL_EVIDENCE_COLUMNS)
    stream = io.BytesIO()
    workbook.save(stream)
    return stream.getvalue()


def scenario_dictionary_csv() -> bytes:
    buffer = io.StringIO()
    writer = csv.writer(buffer, lineterminator="\n")
    writer.writerow(["campo", "unidad", "descripcion"])
    for field, _, unit, description in SCENARIO_COLUMNS:
        writer.writerow([field, unit, description])
    return buffer.getvalue().encode("utf-8-sig")


def history_dictionary_csv() -> bytes:
    buffer = io.StringIO()
    writer = csv.writer(buffer, lineterminator="\n")
    writer.writerow(["campo", "unidad", "descripcion"])
    for field, _, unit, description in HISTORY_COLUMNS:
        writer.writerow([field, unit, description])
    return buffer.getvalue().encode("utf-8-sig")


def _style_data_sheet(sheet, column_count: int) -> None:
    header_fill = PatternFill("solid", fgColor="6F1739")
    for cell in sheet[1]:
        cell.font = Font(bold=True, color="FFFFFF", name="Inter")
        cell.fill = header_fill
        cell.alignment = Alignment(vertical="center", wrap_text=True)
    sheet.row_dimensions[1].height = 34
    sheet.freeze_panes = "A2"
    sheet.auto_filter.ref = f"A1:{get_column_letter(column_count)}1"
    widths = {
        "territory_id": 40,
        "scenario_id": 40,
        "municipality_name": 30,
        "name": 30,
        "input_sha256": 68,
        "model_sha256": 68,
        "sha256": 68,
    }
    for index in range(1, column_count + 1):
        header = str(sheet.cell(row=1, column=index).value)
        sheet.column_dimensions[get_column_letter(index)].width = widths.get(header, 18)


def _write_metadata_sheet(sheet, meta: DatasetMeta, sources: dict | None) -> None:
    sheet.append(["Campo", "Valor"])
    rows: Iterable[tuple[str, Any]] = (
        ("reference_date", meta.reference_date),
        ("model_run_id", str(meta.model_run_id) if meta.model_run_id else ""),
        (
            "dataset_version_ids",
            ", ".join(str(item) for item in meta.dataset_version_ids),
        ),
        ("generated_at", meta.generated_at.isoformat()),
        ("geographic_warning", meta.geographic_warning),
        ("filters", _json_value(meta.filters)),
        ("report_rules_version", REPORT_RULES_VERSION),
        ("sources", _json_value(sources or {})),
        (
            "uso",
            "Agregados municipales; prohibida la inferencia individual. "
            "Los escenarios son hipotéticos y no implican efecto causal.",
        ),
    )
    for key, value in rows:
        sheet.append([key, _escape_spreadsheet(value)])
    sheet.append([])
    sheet.append(["Fuente", "Detalle"])
    for source in (sources or {}).get("sources", []) if isinstance(sources, dict) else []:
        sheet.append(
            [
                _escape_spreadsheet(source.get("year")),
                _escape_spreadsheet(source.get("source_url", "")),
            ]
        )
    for cell in sheet[1]:
        cell.font = Font(bold=True, name="Inter")
    sheet.column_dimensions["A"].width = 30
    sheet.column_dimensions["B"].width = 110


def _write_dictionary_sheet(sheet, columns: Sequence[tuple[str, str, str, str]]) -> None:
    # Tuple order is (campo, etiqueta, tipo/unidad, descripción); keep the
    # header exactly aligned with that order so labels are not silently written
    # under "tipo/unidad".
    sheet.append(["campo", "etiqueta", "tipo/unidad", "descripcion"])
    for field, label, unit, description in columns:
        sheet.append([field, label, unit, description])
    for cell in sheet[1]:
        cell.font = Font(bold=True, name="Inter")
    for index, width in {"A": 32, "B": 22, "C": 90, "D": 52}.items():
        sheet.column_dimensions[index].width = width
    sheet.freeze_panes = "A2"


__all__ = [
    "HISTORY_COLUMNS",
    "HISTORY_FIELDS",
    "MODEL_EVIDENCE_COLUMNS",
    "MODEL_EVIDENCE_FIELDS",
    "REPORT_DICTIONARY",
    "SCENARIO_COLUMNS",
    "SCENARIO_FIELDS",
    "export_history",
    "export_model_evidence",
    "export_scenario_csv",
    "history_dictionary_csv",
    "scenario_dictionary_csv",
]
